"""Read-only, row-streaming .xlsx importer; no full-workbook pandas copy."""
from dataclasses import asdict
from datetime import date, datetime
import hashlib
from zipfile import BadZipFile
from openpyxl.utils.exceptions import InvalidFileException
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from .recording import COLUMNS, DatasetProfile, Observation, Recording, numeric, raw_value


class TimestampParser:
    def __init__(self, profile):
        self.profile = profile
        self.origin = None
        self.awareness = None

    def parse(self, value):
        if self.profile.time_unit != 'datetime':
            value = numeric(value)
            if value is None:
                return None, 'missing_or_nonnumeric_timestamp'
            result = value / {'s': 1, 'ms': 1000, 'us': 1000000}[self.profile.time_unit]
            return (result, None) if result >= 0 else (None, 'negative_timestamp')
        if not isinstance(value, (str, datetime, date)) or not str(value).strip():
            return None, 'datetime_required_numeric_time_unit_unverified'
        try:
            # ISO-8601 or native Excel datetime; no locale/date-format guessing.
            if isinstance(value, str):
                value = datetime.fromisoformat(value.replace('Z', '+00:00'))
            stamp = pd.Timestamp(value)
            if pd.isna(stamp):
                raise ValueError('NaT')
            aware = stamp.tzinfo is not None
            if self.awareness is not None and self.awareness != aware:
                return None, 'mixed_timezone_awareness'
            if self.profile.time_basis == 'unix':
                if not aware:
                    return None, 'timezone_required_for_unix_timestamp'
                result = stamp.timestamp()
            else:
                # Naive Excel dates become relative deltas; never invent a timezone.
                if self.origin is None:
                    self.origin = stamp
                result = (stamp - self.origin).total_seconds()
            self.awareness = aware
            if result < 0:
                # Preserve negative offsets for quality/replay checks; original time remains unchanged.
                return result, 'before_recording_origin'
            return result, None
        except (ValueError, TypeError, OverflowError):
            return None, 'invalid_datetime'


def load_xlsx(source, *, profile=None, sheet=None, max_rows=250_000, filename=None):
    """Return a Recording, preserving invalid rows. Path or seekable binary stream.

    Formula expressions are retained (never evaluated); their signals are flagged
    nonnumeric. Max rows is a configurable safety limit, never silent truncation.
    """
    profile = profile or DatasetProfile()
    if not isinstance(max_rows, int) or max_rows < 1:
        raise ValueError('max_rows must be a positive integer')
    own_stream = isinstance(source, (str, Path))
    stream = open(source, 'rb') if own_stream else source
    if not hasattr(stream, 'seek'):
        raise ValueError('Excel source must be a seekable binary stream')
    position = stream.tell()
    workbook = None
    try:
        stream.seek(0)
        digest = hashlib.sha256()
        while block := stream.read(1024 * 1024):
            digest.update(block)
        stream.seek(0)
        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        if not workbook.sheetnames:
            raise ValueError('workbook contains no sheets')
        if sheet is not None and sheet not in workbook.sheetnames:
            raise ValueError(f'unknown worksheet: {sheet}')
        worksheet = workbook[sheet or workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError('worksheet is empty')
        # Ignore only trailing empty header cells; reject ambiguity inside the table.
        header = list(header)
        while header and header[-1] is None:
            header.pop()
        if not header or any(not isinstance(x, str) or not x.strip() for x in header) or len(set(header)) != len(header):
            raise ValueError('worksheet needs unique, nonempty string column names')
        if 'time' not in header:
            warnings_list = ['Missing time column; timestamps/replay unavailable.']
        else:
            warnings_list = []
        missing = [column for column in COLUMNS if column not in header]
        if missing:
            warnings_list.append('Missing source columns: ' + ', '.join(missing))
        warnings_list.append('Column meanings follow the supplied task schema; dataset documentation and edition are unverified.')
        undeclared = [column for column, (_, unit) in COLUMNS.items() if unit is not None and column in header and column not in profile.units]
        if undeclared:
            warnings_list.append('Source units unverified: ' + ', '.join(undeclared))
        if profile.current_convention == 'unknown':
            warnings_list.append('Current polarity unverified: signed electrical power has no charging/discharging labels.')
        parser = TimestampParser(profile)
        observations = []
        for row_number, cells in enumerate(rows, 2):
            # Entirely blank rows are observations too: preserve recording order and flag them.
            if len(observations) >= max_rows:
                raise ValueError(f'workbook exceeds max_rows={max_rows}; increase the explicit limit or split the workbook')
            if any(value is not None for value in cells[len(header):]):
                raise ValueError(f'row {row_number}: data under unnamed columns')
            raw = {name: raw_value(cells[i] if i < len(cells) else None) for i, name in enumerate(header)}
            values = {target: profile.normalize(column, raw.get(column)) for column, (target, _) in COLUMNS.items()}
            timestamp, reason = parser.parse(cells[header.index('time')] if 'time' in header else None)
            observation = Observation(row_number, timestamp, raw, values)
            if reason:
                observation.flag('timestamp_s', 'TIMESTAMP_ERROR', 'import:' + reason)
            observations.append(observation)
        if not observations:
            raise ValueError('worksheet has no observations')
        signals = {target: {
            'source_column': column, 'source_unit': profile.units.get(column),
            'canonical_unit': canonical, 'unit_status': 'user_declared' if column in profile.units else 'unverified',
            'meaning_status': 'task_schema_unverified',
            'present': column in header,
        } for column, (target, canonical) in COLUMNS.items()}
        metadata = {
            'schema': 'can-lab-recording-1', 'data_kind': 'imported_decoded_telemetry',
            'source_filename': Path(filename or getattr(stream, 'name', 'uploaded.xlsx')).name,
            'sha256': digest.hexdigest(), 'sheet': worksheet.title,
            'available_sheets': workbook.sheetnames, 'source_columns': header,
            'profile': asdict(profile), 'signals': signals,
            'normalized_time_basis': profile.time_basis,
            'time_origin': parser.origin.isoformat() if parser.origin is not None else None,
            'warnings': warnings_list,
        }
        return Recording(metadata, observations)
    except (BadZipFile, InvalidFileException) as error:
        raise ValueError('input is not a readable .xlsx workbook') from error
    finally:
        if workbook is not None:
            workbook.close()
        if own_stream:
            stream.close()
        else:
            stream.seek(position)
